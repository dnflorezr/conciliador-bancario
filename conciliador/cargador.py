"""Lectura de extractos y auxiliares desde CSV o Excel.

Ningún banco entrega el mismo formato, así que el cargador no asume nombres de
columna: recibe un mapeo y normaliza lo que encuentre. Las dos conversiones que
más fallan en la práctica son las fechas (dd/mm/aaaa contra aaaa-mm-dd) y los
importes (separador de miles con punto y decimal con coma, paréntesis para
negativos, símbolo de moneda pegado al número).
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from .modelos import Movimiento, Origen

FORMATOS_FECHA = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%m/%d/%Y",
    "%Y/%m/%d",
    "%d.%m.%Y",
)


class ErrorDeCarga(Exception):
    """Un archivo no se pudo interpretar."""


@dataclass
class MapeoColumnas:
    """Qué columna del archivo corresponde a cada campo."""

    fecha: str = "fecha"
    importe: str = "importe"
    descripcion: str | None = "descripcion"
    referencia: str | None = "referencia"
    id: str | None = None


def parsear_fecha(valor: object) -> date:
    """Convierte a fecha probando los formatos habituales."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    if not texto:
        raise ErrorDeCarga("Fecha vacía.")
    for fmt in FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    raise ErrorDeCarga(f"No se reconoce el formato de fecha: {texto!r}")


def parsear_importe(valor: object) -> Decimal:
    """Convierte a Decimal tolerando los formatos de monto más comunes.

    Acepta ``1.234,56`` y ``1,234.56``, el símbolo de moneda, los espacios y
    los paréntesis contables para negativos.
    """
    if isinstance(valor, Decimal):
        return valor
    if isinstance(valor, int):
        return Decimal(valor)
    if isinstance(valor, float):
        # Pasa por str para no arrastrar el error binario del float.
        return Decimal(str(valor))

    texto = str(valor).strip()
    if not texto:
        raise ErrorDeCarga("Importe vacío.")

    negativo = texto.startswith("(") and texto.endswith(")")
    if negativo:
        texto = texto[1:-1]

    texto = re.sub(r"[^\d,.\-]", "", texto)

    if "," in texto and "." in texto:
        # El separador decimal es el que aparece más a la derecha.
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        entero, _, dec = texto.rpartition(",")
        texto = f"{entero.replace(',', '')}.{dec}" if len(dec) in (1, 2) else texto.replace(",", "")

    try:
        monto = Decimal(texto)
    except InvalidOperation as exc:
        raise ErrorDeCarga(f"No se reconoce el importe: {valor!r}") from exc

    return -monto if negativo else monto


def _filas_csv(ruta: Path) -> Iterable[dict[str, str]]:
    with ruta.open(encoding="utf-8-sig", newline="") as fh:
        muestra = fh.read(4096)
        fh.seek(0)
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
        except csv.Error:
            dialecto = csv.excel
        yield from csv.DictReader(fh, dialect=dialecto)


def _filas_excel(ruta: Path) -> Iterable[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ErrorDeCarga("Para leer Excel hace falta openpyxl.") from exc

    libro = load_workbook(ruta, read_only=True, data_only=True)
    hoja = libro.active
    filas = hoja.iter_rows(values_only=True)
    try:
        encabezado = [str(c).strip() if c is not None else "" for c in next(filas)]
    except StopIteration:
        return
    for fila in filas:
        if all(c is None for c in fila):
            continue
        yield dict(zip(encabezado, fila))
    libro.close()


def cargar(
    ruta: str | Path,
    origen: Origen,
    mapeo: MapeoColumnas | None = None,
    prefijo_id: str | None = None,
) -> list[Movimiento]:
    """Lee un archivo y devuelve movimientos listos para conciliar.

    Args:
        ruta: Archivo ``.csv``, ``.xlsx`` o ``.xlsm``.
        origen: Si las filas son del banco o del libro.
        mapeo: Nombres de columna del archivo. Por defecto usa los canónicos.
        prefijo_id: Prefijo de los ids generados cuando el archivo no trae uno.

    Raises:
        ErrorDeCarga: Si falta una columna obligatoria o una fila no se puede
            interpretar. El mensaje incluye el número de fila.
    """
    ruta = Path(ruta)
    if not ruta.exists():
        raise ErrorDeCarga(f"No existe el archivo {ruta}")

    mapeo = mapeo or MapeoColumnas()
    prefijo = prefijo_id or ("B" if origen is Origen.BANCO else "L")

    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        filas = _filas_excel(ruta)
    else:
        filas = _filas_csv(ruta)

    movimientos: list[Movimiento] = []
    for n, fila in enumerate(filas, start=2):
        normal = {str(k).strip(): v for k, v in fila.items() if k is not None}
        if mapeo.fecha not in normal or mapeo.importe not in normal:
            raise ErrorDeCarga(
                f"Fila {n}: faltan las columnas '{mapeo.fecha}' o '{mapeo.importe}'. "
                f"Columnas encontradas: {sorted(normal)}"
            )
        try:
            fecha = parsear_fecha(normal[mapeo.fecha])
            importe = parsear_importe(normal[mapeo.importe])
        except ErrorDeCarga as exc:
            raise ErrorDeCarga(f"Fila {n} de {ruta.name}: {exc}") from exc

        ident = None
        if mapeo.id and normal.get(mapeo.id) not in (None, ""):
            ident = str(normal[mapeo.id]).strip()

        def _texto(col: str | None) -> str | None:
            if not col:
                return None
            v = normal.get(col)
            return str(v).strip() if v not in (None, "") else None

        movimientos.append(
            Movimiento(
                id=ident or f"{prefijo}{n - 1:04d}",
                fecha=fecha,
                importe=importe,
                descripcion=_texto(mapeo.descripcion) or "",
                referencia=_texto(mapeo.referencia),
                origen=origen,
            )
        )

    if not movimientos:
        raise ErrorDeCarga(f"{ruta.name} no tiene filas de datos.")
    return movimientos
